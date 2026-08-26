"""Creating real Word/Excel files — via python-docx/openpyxl, both already
installed in this project's venv (verified live, not assumed).

Deliberately simple: a title, paragraphs, and an optional bulleted list for
documents; a single sheet of rows for spreadsheets. Wisp is a local assistant
generating something the user asked for on the spot, not a document-authoring
tool — richer formatting (multi-sheet workbooks, styled templates, tables of
contents) is exactly what the user's own Word/Excel is for once the file
exists.
"""
from __future__ import annotations

from pathlib import Path

from service.tools.registry import register


@register(
    "write_document",
    "Create a real Word (.docx) document with a title, paragraphs, and an "
    "optional bulleted list. Use for 'write this up as a document' or "
    "'make me a Word doc of this' — not write_file, which only produces "
    "plain text.",
    {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Where to save it, e.g. '~/Documents/Notes.docx'."},
            "title": {"type": "string", "description": "Document title, as a heading."},
            "paragraphs": {"type": "array", "items": {"type": "string"},
                           "description": "Body paragraphs, in order."},
            "bullets": {"type": "array", "items": {"type": "string"},
                       "description": "Optional bulleted list, appended after the paragraphs."},
        },
        "required": ["path", "title"],
    },
    category="fs_write",
    aliases=["make me a word document of this", "write this up as a docx",
             "create a document with this content", "save this as a word file"],
)
def write_document(path: str, title: str, paragraphs: list[str] | None = None,
                   bullets: list[str] | None = None) -> str:
    from service.tools.builtin import _wrong_account_path

    if (msg := _wrong_account_path(path)):
        return msg
    import docx

    p = Path(path).expanduser()
    if p.suffix.lower() != ".docx":
        p = p.with_suffix(".docx")
    if p.exists():
        return f"(refusing to overwrite {p} — it already exists. Pick a different name.)"

    doc = docx.Document()
    doc.add_heading(title.strip() or "Untitled", level=1)
    for para in (paragraphs or []):
        if para.strip():
            doc.add_paragraph(para)
    for item in (bullets or []):
        if item.strip():
            doc.add_paragraph(item, style="List Bullet")
    try:
        p.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(p))
    except Exception as e:  # noqa: BLE001
        return f"(error saving document: {e})"
    return f"Created {p.name} ({p.stat().st_size // 1024}KB) at {p.parent}."


@register(
    "spreadsheet_ops",
    "Create a real Excel (.xlsx) spreadsheet with a header row and data "
    "rows. Use for 'make me a spreadsheet of this' — not write_file, which "
    "only produces plain text/CSV.",
    {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Where to save it, e.g. '~/Documents/Budget.xlsx'."},
            "headers": {"type": "array", "items": {"type": "string"}, "description": "Column headers."},
            "rows": {"type": "array", "items": {"type": "array"},
                     "description": "Data rows — each an array of cell values matching the headers."},
        },
        "required": ["path", "headers"],
    },
    category="fs_write",
    aliases=["make me a spreadsheet of this", "save this as an excel file",
             "create an xlsx with these columns", "put this data in a spreadsheet"],
)
def spreadsheet_ops(path: str, headers: list[str], rows: list[list] | None = None) -> str:
    from service.tools.builtin import _wrong_account_path

    if (msg := _wrong_account_path(path)):
        return msg
    import openpyxl

    p = Path(path).expanduser()
    if p.suffix.lower() != ".xlsx":
        p = p.with_suffix(".xlsx")
    if p.exists():
        return f"(refusing to overwrite {p} — it already exists. Pick a different name.)"
    if not headers:
        return "(error: spreadsheet_ops needs at least one header.)"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    for row in (rows or []):
        ws.append(list(row))
    # Column widths sized to content, capped — an unbounded width from one
    # long cell would otherwise make the sheet awkward to open.
    for col in ws.columns:
        width = min(60, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2)
        ws.column_dimensions[col[0].column_letter].width = width

    try:
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(p))
    except Exception as e:  # noqa: BLE001
        return f"(error saving spreadsheet: {e})"
    return f"Created {p.name} ({len(rows or [])} row(s)) at {p.parent}."
